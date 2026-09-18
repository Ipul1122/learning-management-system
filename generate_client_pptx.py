import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_client_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions (Design System LMS)
    COLOR_BG_LIGHT = RGBColor(250, 248, 245)      # #FAF8F5 Warm Light Canvas
    COLOR_BG_WHITE = RGBColor(255, 255, 255)      # #FFFFFF Card Surface
    COLOR_BG_DARK = RGBColor(30, 27, 24)          # #1E1B18 Deep Charcoal
    COLOR_CARD_DARK = RGBColor(42, 38, 34)        # #2A2622
    COLOR_PRIMARY_ORANGE = RGBColor(255, 107, 0)  # #FF6B00 Signature Brand
    COLOR_ACCENT_ORANGE = RGBColor(249, 115, 22)  # #F97316
    COLOR_CRIMSON = RGBColor(225, 29, 72)         # #E11D48 Sub-Main Red
    COLOR_EMERALD = RGBColor(16, 185, 129)        # #10B981 Success
    COLOR_SKY = RGBColor(14, 165, 233)            # #0EA5E9 Info/Zoom
    COLOR_TEXT_MAIN = RGBColor(30, 27, 24)        # #1E1B18 Text
    COLOR_TEXT_MUTED = RGBColor(110, 103, 95)     # #6E675F Secondary
    COLOR_TEXT_WHITE = RGBColor(255, 255, 255)
    COLOR_BORDER_LIGHT = RGBColor(235, 229, 223)  # #EBE5DF
    COLOR_BORDER_DARK = RGBColor(60, 54, 49)      # #3C3631
    COLOR_ORANGE_LIGHT = RGBColor(255, 247, 237)  # #FFF7ED

    screenshots_dir = os.path.abspath("screenshots")

    def set_slide_background(slide, color):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = color
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, category, title, is_dark=False):
        # Category Chip
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.6), Inches(0.32))
        chip.fill.solid()
        chip.fill.fore_color.rgb = COLOR_CARD_DARK if is_dark else COLOR_ORANGE_LIGHT
        chip.line.color.rgb = COLOR_PRIMARY_ORANGE if is_dark else RGBColor(254, 215, 170)
        chip.line.width = Pt(1)
        tf_chip = chip.text_frame
        tf_chip.word_wrap = True
        tf_chip.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_chip = tf_chip.paragraphs[0]
        p_chip.text = "  ●  " + category.upper()
        p_chip.font.name = "Montserrat"
        p_chip.font.size = Pt(9.5)
        p_chip.font.bold = True
        p_chip.font.color.rgb = COLOR_PRIMARY_ORANGE

        # Title
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Montserrat"
        p.font.size = Pt(21)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_WHITE if is_dark else COLOR_TEXT_MAIN

    def create_card(slide, left, top, width, height, bg_color=COLOR_BG_WHITE, border_color=COLOR_BORDER_LIGHT):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    def add_feature_slide_with_image(category, title, subtitle, bullets, image_filename, image_caption):
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide, COLOR_BG_LIGHT)
        add_header(slide, category, title)

        # Left Column: Narrative Card
        left_card = create_card(slide, Inches(0.8), Inches(1.5), Inches(5.3), Inches(5.5))
        tfl = left_card.text_frame
        tfl.margin_left = Inches(0.35); tfl.margin_right = Inches(0.35); tfl.margin_top = Inches(0.3)
        tfl.word_wrap = True

        p_sub = tfl.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = "Montserrat"; p_sub.font.size = Pt(13); p_sub.font.bold = True; p_sub.font.color.rgb = COLOR_PRIMARY_ORANGE
        p_sub.space_after = Pt(12)

        for b_title, b_desc in bullets:
            pb = tfl.add_paragraph()
            pb.text = "✓  " + b_title
            pb.font.name = "Montserrat"; pb.font.size = Pt(11); pb.font.bold = True; pb.font.color.rgb = COLOR_TEXT_MAIN
            pb.space_before = Pt(6)

            pbd = tfl.add_paragraph()
            pbd.text = b_desc
            pbd.font.name = "Quicksand"; pbd.font.size = Pt(10); pbd.font.color.rgb = COLOR_TEXT_MUTED
            pbd.space_after = Pt(8)

        # Right Column: Visual Frame Card + Embedded Screenshot
        right_card = create_card(slide, Inches(6.3), Inches(1.5), Inches(6.2), Inches(5.5))
        
        img_path = os.path.join(screenshots_dir, image_filename)
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(6.45), Inches(1.65), Inches(5.9), Inches(3.68))
        
        # Caption below screenshot
        cap_box = slide.shapes.add_textbox(Inches(6.45), Inches(5.45), Inches(5.9), Inches(1.4))
        tfc = cap_box.text_frame
        tfc.word_wrap = True
        tfc.margin_left = Inches(0.1); tfc.margin_right = Inches(0.1); tfc.margin_top = Inches(0.05)
        
        pc0 = tfc.paragraphs[0]
        pc0.text = "TAMPILAN ANTARMUKA AKTIF (LIVE SYSTEM PREVIEW):"
        pc0.font.name = "Montserrat"; pc0.font.size = Pt(8.5); pc0.font.bold = True; pc0.font.color.rgb = COLOR_PRIMARY_ORANGE
        pc0.space_after = Pt(2)

        pc1 = tfc.add_paragraph()
        pc1.text = image_caption
        pc1.font.name = "Quicksand"; pc1.font.size = Pt(9.5); pc1.font.color.rgb = COLOR_TEXT_MUTED

        return slide

    # =========================================================================
    # SLIDE 1: COVER SLIDE (CLIENT EXECUTIVE THEME)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, COLOR_BG_DARK)

    # Accent glow top
    glow = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.12))
    glow.fill.solid(); glow.fill.fore_color.rgb = COLOR_PRIMARY_ORANGE; glow.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(3.4))
    tf1 = tb1.text_frame; tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SOLUSI DIGITALISASI PELATIHAN & SERTIFIKASI MULTI-CABANG"
    p.font.name = "Montserrat"; p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_ORANGE
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "Platform Learning Management System\n(LMS) Terpadu Multi-Cabang"
    p.font.name = "Montserrat"; p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = COLOR_TEXT_WHITE
    p.space_after = Pt(16)

    p = tf1.add_paragraph()
    p.text = "Presentasi Eksekutif: Alur Bisnis, Tata Kelola Pengguna, Standar UI/UX, dan Panduan Operasional Pelatihan 20 Jam Pelajaran (JP)"
    p.font.name = "Quicksand"; p.font.size = Pt(15); p.font.color.rgb = RGBColor(214, 211, 209)

    # Stat Highlights Bottom
    exec_stats = [
        ("Multi-Cabang Terpadu", "Sentralisasi kendali pusat dengan otonomi operasional cabang daerah"),
        ("Kepatuhan 20 JP", "Perhitungan otomatis 900 menit belajar terstandarisasi kurikulum"),
        ("Proteksi Kuota 40", "Jaminan batas maksimal 40 kursi kelas fisik bebas overbooking"),
        ("E-Sertifikat QR", "Otentikasi sertifikat digital resmi yang dapat divalidasi publik")
    ]
    card_w = Inches(2.7)
    card_h = Inches(1.25)
    for i, (head, desc) in enumerate(exec_stats):
        cx = Inches(1.0) + i * (card_w + Inches(0.17))
        c = create_card(slide1, cx, Inches(5.4), card_w, card_h, COLOR_CARD_DARK, COLOR_BORDER_DARK)
        tfc = c.text_frame; tfc.word_wrap = True; tfc.margin_left = Inches(0.2); tfc.margin_top = Inches(0.15)
        p1 = tfc.paragraphs[0]; p1.text = head; p1.font.name = "Montserrat"; p1.font.size = Pt(12.5); p1.font.bold = True; p1.font.color.rgb = COLOR_PRIMARY_ORANGE
        p2 = tfc.add_paragraph(); p2.text = desc; p2.font.name = "Quicksand"; p2.font.size = Pt(9.5); p2.font.color.rgb = RGBColor(168, 162, 158)

    # =========================================================================
    # SLIDE 2: LATAR BELAKANG & NILAI TAMBAH BISNIS (CLIENT VALUE)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, COLOR_BG_LIGHT)
    add_header(slide2, "Nilai Bisnis", "Mengapa Platform LMS Ini Penting Bagi Transformasi Lembaga?")

    col_w = Inches(3.75); col_h = Inches(5.3); top_pos = Inches(1.6)

    c1 = create_card(slide2, Inches(0.8), top_pos, col_w, col_h)
    tf1 = c1.text_frame; tf1.margin_left = Inches(0.3); tf1.margin_top = Inches(0.3)
    p = tf1.paragraphs[0]; p.text = "Tantangan Sebelum Sistem"; p.font.name = "Montserrat"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_CRIMSON; p.space_after = Pt(12)
    b1 = [
        "Pengelolaan Terpisah-Pisah: Tiap cabang mengelola pelatihan dengan cara manual, data siswa tercecer.",
        "Overbooking Kelas Fisik: Sulit mengunci kuota 40 kursi saat pendaftaran dibuka serentak.",
        "Sertifikat Rentan Pemalsuan: Penggunaan sertifikat cetak/PDF biasa tanpa kode verifikasi keabsahan.",
        "Pencatatan Jam Belajar Bias: Kesulitan membuktikan siswa telah menyelesaikan minimal 20 JP (900 menit)."
    ]
    for b in b1:
        pb = tf1.add_paragraph(); pb.text = "✕ " + b; pb.font.name = "Quicksand"; pb.font.size = Pt(10.5); pb.font.color.rgb = COLOR_TEXT_MUTED; pb.space_after = Pt(8)

    c2 = create_card(slide2, Inches(4.78), top_pos, col_w, col_h)
    tf2 = c2.text_frame; tf2.margin_left = Inches(0.3); tf2.margin_top = Inches(0.3)
    p = tf2.paragraphs[0]; p.text = "Solusi Inovatif LMS"; p.font.name = "Montserrat"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_ORANGE; p.space_after = Pt(12)
    b2 = [
        "Sentralisasi Kendali Multi-Cabang: Satu portal tunggal untuk mengawasi operasional seluruh unit cabang.",
        "Kunci Kuota Realtime: Otomatis menutup pendaftaran tatap muka tepat saat kuota ke-40 terisi.",
        "QR Code Validasi Publik: Lembaga eksternal / perusahaan dapat mengecek keaslian dokumen secara instan.",
        "Engine 20 JP Terstandar: Jam pembelajaran tercatat otomatis dari presensi sesi dan 1-klik masuk Zoom."
    ]
    for b in b2:
        pb = tf2.add_paragraph(); pb.text = "✓ " + b; pb.font.name = "Quicksand"; pb.font.size = Pt(10.5); pb.font.color.rgb = COLOR_TEXT_MUTED; pb.space_after = Pt(8)

    c3 = create_card(slide2, Inches(8.76), top_pos, col_w, col_h)
    tf3 = c3.text_frame; tf3.margin_left = Inches(0.3); tf3.margin_right = Inches(0.3); tf3.margin_top = Inches(0.3)
    p = tf3.paragraphs[0]; p.text = "Keuntungan Bagi Klien"; p.font.name = "Montserrat"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_EMERALD; p.space_after = Pt(12)
    b3 = [
        "Peningkatan Reputasi & Akreditasi: Menjamin lulusan pelatihan benar-benar memenuhi standar sertifikasi kejuruan.",
        "Efisiensi Biaya & Waktu: Memangkas beban kerja administrasi cabang hingga lebih dari 70%.",
        "Pengalaman Pengguna Modern: Desain antarmuka bersih dan ramah pengguna yang disukai peserta & instruktur.",
        "Audit Trail Lengkap: Setiap aktivitas penerbitan sertifikat dan mutasi data terekam transparan."
    ]
    for b in b3:
        pb = tf3.add_paragraph(); pb.text = "★ " + b; pb.font.name = "Quicksand"; pb.font.size = Pt(10.5); pb.font.color.rgb = COLOR_TEXT_MUTED; pb.space_after = Pt(8)

    # =========================================================================
    # SLIDE 3: HIERARKI 4 PERAN PENGGUNA (ROLE GOVERNANCE)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, COLOR_BG_LIGHT)
    add_header(slide3, "Tata Kelola Peran", "Pemisahan Hak Akses & Tugas Pengguna (Multi-Role Governance)")

    roles = [
        ("SUPER ADMIN", "Direksi & Pengawas Pusat", COLOR_CRIMSON, [
            "Membuat & mengontrol master cabang se-Indonesia",
            "Menugaskan Admin Cabang pengelola wilayah",
            "Memantau seluruh aktivitas sistem skala nasional",
            "Menerbitkan berita & pengumuman global lembaga"
        ]),
        ("ADMIN CABANG", "Manajer Operasional Cabang", COLOR_PRIMARY_ORANGE, [
            "Mengelola akun Trainer & penugasan mengajar",
            "Membuka kelas pelatihan (Offline, Online, Hybrid)",
            "Menentukan jadwal sesi pertemuan & link awal Zoom",
            "Mengawasi ketersediaan bank soal ujian cabang"
        ]),
        ("TRAINER", "Instruktur Profesional", COLOR_SKY, [
            "Menyusun paket kuis & ujian evaluasi kompetensi",
            "Mengaktifkan live meeting Zoom sesi pembelajaran",
            "Melakukan audit pemenuhan 20 JP & nilai kuis",
            "Memutuskan kelulusan sertifikat atau jadwal remedial"
        ]),
        ("PESERTA", "Siswa / Trainee Pelatihan", COLOR_EMERALD, [
            "Mendaftar akun aman via verifikasi kode email OTP",
            "Mengeksplorasi katalog & memilih kelas pelatihan",
            "Menghadiri sesi (1-Klik Zoom) & memantau progres 20 JP",
            "Menyelesaikan evaluasi & mengunduh sertifikat resmi"
        ])
    ]
    card_w = Inches(2.78); card_h = Inches(5.3)
    for i, (r_name, r_sub, r_col, r_perms) in enumerate(roles):
        cx = Inches(0.8) + i * Inches(2.98)
        c = create_card(slide3, cx, Inches(1.6), card_w, card_h)
        tfc = c.text_frame; tfc.margin_left = Inches(0.2); tfc.margin_top = Inches(0.25)
        p0 = tfc.paragraphs[0]; p0.text = r_name; p0.font.name = "Montserrat"; p0.font.size = Pt(13.5); p0.font.bold = True; p0.font.color.rgb = r_col
        p1 = tfc.add_paragraph(); p1.text = r_sub; p1.font.name = "Quicksand"; p1.font.size = Pt(10); p1.font.color.rgb = COLOR_TEXT_MUTED; p1.space_after = Pt(14)
        p_div = tfc.add_paragraph(); p_div.text = "Fungsi Utama:"; p_div.font.name = "Montserrat"; p_div.font.size = Pt(9.5); p_div.font.bold = True; p_div.font.color.rgb = COLOR_TEXT_MAIN; p_div.space_after = Pt(6)
        for p_item in r_perms:
            pi = tfc.add_paragraph(); pi.text = "✓ " + p_item; pi.font.name = "Quicksand"; pi.font.size = Pt(9.5); pi.font.color.rgb = COLOR_TEXT_MUTED; pi.space_after = Pt(6)

    # =========================================================================
    # SLIDE 4: SUPER ADMIN DASHBOARD (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Pusat Kendali",
        title="Dashboard Super Admin: Supervisi Operasional Seluruh Cabang",
        subtitle="Kendali Penuh Skala Nasional Dalam Satu Dasbor Eksekutif",
        bullets=[
            ("Statistik Realtime Cabang", "Menampilkan metrik komprehensif total cabang aktif, jumlah akun admin, total tenaga pengajar (trainer), dan akumulasi siswa terdaftar."),
            ("Manajemen Kantor Cabang", "Kemudahan menambah kantor cabang baru, memperbarui profil alamat, serta menonaktifkan unit tanpa kehilangan data historis."),
            ("Audit Trail & Log Aktivitas", "Memantau setiap aksi penting yang dilakukan oleh seluruh admin cabang di Indonesia untuk mencegah penyalahgunaan wewenang."),
            ("Publikasi Informasi Nasional", "Menerbitkan pengumuman serentak yang langsung tampil pada halaman beranda seluruh siswa di semua cabang.")
        ],
        image_filename="super_admin.png",
        image_caption="Tangkapan layar nyata dashboard Super Admin dengan modul operasional, metrik statistik 4 entitas, dan navigasi terpadu."
    )

    # =========================================================================
    # SLIDE 5: ADMIN CABANG DASHBOARD (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Operasional Cabang",
        title="Dashboard Admin Cabang: Manajemen Kelas & Tenaga Pendidik",
        subtitle="Otonomi Pengelolaan Pelatihan di Tingkat Wilayah Daerah",
        bullets=[
            ("Distribusi Trainer Cabang", "Membuat akun instruktur lokal, mengalokasikan pengajar ke kelas tertentu, dan memantau beban jadwal mengajar."),
            ("Penerbitan Kelas Pelatihan", "Mengatur kelas baru dengan pilihan mode Offline, Online, atau Hybrid serta menetapkan target jam pelajaran."),
            ("Penjadwalan Sesi & Zoom", "Menyusun tanggal pertemuan, durasi menit tiap sesi pembelajaran, dan menyematkan tautan rapat online Zoom."),
            ("Pemantauan Kuota Peserta", "Melihat perkembangan keterisian kursi kelas secara langsung guna mengantisipasi kelas yang hampir penuh.")
        ],
        image_filename="admin_cabang.png",
        image_caption="Tangkapan layar nyata dashboard Admin Cabang menampilkan metrik lokal, daftar kelas berjalan, dan modul manajemen instruktur."
    )

    # =========================================================================
    # SLIDE 6: TRAINER DASHBOARD (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Portal Instruktur",
        title="Dashboard Trainer: Pusat Kurikulum, Kelas & Live Teaching",
        subtitle="Ruang Kerja Terpadu Instruktur Mengawal Kualitas Kelulusan",
        bullets=[
            ("Akses Cepat Pengajaran", "Trainer langsung melihat kelas aktif yang diampunya, jadwal mengajar hari ini, dan jumlah siswa yang diajar."),
            ("Monitoring Siswa Siap Uji", "Notifikasi otomatis saat ada peserta yang berhasil menuntaskan jam belajar dan siap untuk dievaluasi kelulusannya."),
            ("Penyusunan Paket Ujian Kuis", "Merancang paket kuis evaluasi berbasis bank soal cabang dengan batas waktu dan passing grade."),
            ("Interaksi Forum Komunitas", "Menjawab pertanyaan materi dan berkonsultasi langsung dengan siswa pada forum diskusi kelas.")
        ],
        image_filename="trainer_dashboard.png",
        image_caption="Tangkapan layar nyata dashboard Trainer dengan metrik pengajaran, jadwal live session terdekat, dan modul akses cepat."
    )

    # =========================================================================
    # SLIDE 7: MANAJEMEN KELAS & 1-KLIK ZOOM TRAINER (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Live Teaching",
        title="Manajemen Kelas & Sesi Zoom: Memulai Pertemuan Tanpa Kendala",
        subtitle="Fleksibilitas Instruktur Menyelenggarakan Tatap Muka Virtual",
        bullets=[
            ("Quick Zoom Session Launcher", "Fitur 1-klik untuk memperbarui tautan Zoom instan tepat sebelum sesi kelas online dimulai."),
            ("Silabus Pertemuan Terstruktur", "Menampilkan rincian agenda per pertemuan, bobot jam pelajaran (JP), dan materi yang akan dibahas."),
            ("Pencatatan Presensi Otomatis", "Siswa yang hadir dan bergabung ke dalam sesi Zoom otomatis terdata waktu belajarnya."),
            ("Status Kesiapan Sesi", "Indikator warna cerdas membedakan sesi yang tautan Zoom-nya sudah siap vs yang belum disiapkan.")
        ],
        image_filename="trainer_classes.png",
        image_caption="Tangkapan layar panel pengelolaan kelas Trainer dengan fitur Quick Zoom terintegrasi dan kartu kelas berdesain bersih."
    )

    # =========================================================================
    # SLIDE 8: KATALOG KELAS & ATURAN KUOTA 40 KURSI (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Katalog Pelatihan",
        title="Katalog Kelas & Aturan Kapasitas: Bebas Overbooking",
        subtitle="Pencegahan Kapasitas Lebih Pada Ruang Kelas Tatap Muka Fisik",
        bullets=[
            ("Proteksi Ketat Kelas Offline (40 Kursi)", "Kapasitas kelas tatap muka dikunci maksimal 40 kursi. Sistem menjamin tidak akan terjadi overbooking meskipun ribuan siswa mendaftar bersamaan."),
            ("Opsi Kelas Online Fleksibel", "Untuk kelas daring, kuota dibuka berskala besar (ratusan peserta) melalui video conference Zoom tanpa batasan kursi."),
            ("Pilihan Jalur Kelas Hybrid", "Siswa dapat memilih ingin hadir langsung di ruang kelas fisik (maks 40) atau mengikuti daring dari rumah."),
            ("Pencarian & Filter Cerdas", "Siswa dapat memfilter katalog berdasarkan cabang kota terdekat, nama materi, atau tipe metode belajar.")
        ],
        image_filename="peserta_catalog.png",
        image_caption="Tangkapan layar katalog kelas peserta dengan indikator sisa kursi realtime (Tersisa X dari 40 Kursi) dan tombol pendaftaran instan."
    )

    # =========================================================================
    # SLIDE 9: DASHBOARD PESERTA & TRACKING 20 JP (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Pengalaman Siswa",
        title="Dashboard Siswa: Pelacakan Progres 20 Jam Pelajaran (JP)",
        subtitle="Transparansi Penuh Perjalanan Belajar Hingga Meraih Sertifikat",
        bullets=[
            ("Standar Akreditasi 20 JP", "Sistem mengacu pada regulasi kurikulum resmi: 1 Jam Pelajaran (JP) = 45 Menit, sehingga target kelulusan adalah akumulasi 900 Menit Belajar."),
            ("Widget Progress Bar Realtime", "Siswa dapat melihat secara visual berapa jam pelajaran yang sudah terkumpul (misal: '16.0 / 20.0 JP (80%) Tuntas')."),
            ("Akses Cepat 3 Langkah", "Kartu navigasi intuitif untuk (1) Menjelajahi Katalog, (2) Masuk Ruang Belajar, dan (3) Mengunduh E-Sertifikat."),
            ("Jadwal Belajar Terdekat", "Pemberitahuan sesi tatap muka atau live Zoom berikutnya agar siswa tidak ketinggalan jadwal kelas.")
        ],
        image_filename="peserta_dashboard.png",
        image_caption="Tangkapan layar nyata dashboard belajar siswa berdesain warm-light dengan banner akumulasi 20 JP dan kartu modul interaktif."
    )

    # =========================================================================
    # SLIDE 10: RUANG BELAJAR & PRESENSI (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Ruang Belajar",
        title="Ruang Belajar Interaktif: Masuk Zoom & Kuis Evaluasi Mandiri",
        subtitle="Semua Kebutuhan Pembelajaran Siswa Terpusat di Satu Tempat",
        bullets=[
            ("Tombol 1-Klik Masuk Zoom", "Siswa tidak perlu lagi mencari tautan meeting di grup chat; cukup menekan satu tombol di silabus untuk langsung tersambung."),
            ("Pengerjaan Kuis Berdurasi Waktu", "Siswa menguji pemahaman modul lewat kuis ber-timer mundur dengan pengacakan butir soal otomatis."),
            ("Pembahasan Kunci Jawaban", "Setelah menyelesaikan kuis, siswa dapat langsung meninjau pembahasan jawaban yang benar beserta penjelasan teori pendukung."),
            ("Pengajuan Kelulusan Mandiri", "Begitu jam belajar genap 20 JP dan seluruh kuis tuntas, sistem otomatis mengajukan berkas ke instruktur.")
        ],
        image_filename="peserta_study.png",
        image_caption="Tangkapan layar halaman ruang belajar siswa dengan tracking pemenuhan jam pelajaran, silabus pertemuan, dan akses materi."
    )

    # =========================================================================
    # SLIDE 11: AUDIT KELULUSAN OLEH TRAINER (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Verifikasi Mutu",
        title="Meja Audit Trainer: Verifikasi Objektif Sebelum Menerbitkan Sertifikat",
        subtitle="Keputusan Kelulusan Berdasarkan Data Kehadiran & Capaian Nilai",
        bullets=[
            ("Audit 3 Aspek Kelulusan", "Trainer memeriksa secara transparan: (1) Akumulasi jam pelajaran (minimal 20 JP), (2) Presensi kehadiran sesi, dan (3) Skor kuis evaluasi."),
            ("Pemberian Keputusan Adil (Approve / Reject)", "Instruktur memiliki wewenang menyetujui kelulusan atau mewajibkan remedial jika kompetensi siswa belum memenuhi batas minimal."),
            ("Instruksi Remedial Tertulis", "Jika ditolak, sistem mewajibkan instruktur mengisi catatan arahan perbaikan materi yang harus dipelajari ulang oleh siswa."),
            ("Otomatisasi Penerbitan", "Ketika tombol 'Setujui' ditekan, sistem secara otomatis menerbitkan dokumen E-Sertifikat resmi dan nomor registrasi.")
        ],
        image_filename="trainer_graduations.png",
        image_caption="Tangkapan layar meja verifikasi kelulusan Trainer dengan tabel antrean peserta, indikator JP tercapai, dan tombol audit performa."
    )

    # =========================================================================
    # SLIDE 12: E-SERTIFIKAT DIGITAL & QR CODE (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Sertifikasi Digital",
        title="E-Sertifikat Resmi & Otentikasi Publik Berbasis QR Code",
        subtitle="Jaminan Keabsahan Dokumen Yang Dapat Diverifikasi Tanpa Login",
        bullets=[
            ("Format PDF Berkualitas Tinggi", "Desain sertifikat elegan siap cetak lengkap dengan nama peserta, judul kejuruan, beban 20 JP, dan tanda tangan digital."),
            ("Nomor Registrasi Nasional Unik", "Setiap sertifikat memiliki kode penomoran unik yang terdaftar resmi pada basis data lembaga."),
            ("QR Code Validasi Publik", "Perusahaan perekrut atau instansi pemerintah cukup memindai QR Code dengan kamera ponsel untuk memastikan keaslian dokumen."),
            ("Halaman Verifikasi Terbuka", "Menampilkan rincian tanggal terbit, nama instruktur pengampu, dan status sertifikat masih aktif berlaku.")
        ],
        image_filename="peserta_certificates.png",
        image_caption="Tangkapan layar portal E-Sertifikat peserta berdesain kartu modern dengan tombol preview PDF dan riwayat sertifikasi."
    )

    # =========================================================================
    # SLIDE 13: FORUM KOMUNITAS & PUSAT INFORMASI (DENGAN SCREENSHOT)
    # =========================================================================
    add_feature_slide_with_image(
        category="Komunikasi & Media",
        title="Pusat Informasi & Forum Komunitas: Kolaborasi Tanpa Batas",
        subtitle="Membangun Budaya Diskusi Aktif dan Keterbukaan Informasi",
        bullets=[
            ("Forum Tanya Jawab Per Kelas", "Siswa dapat membuat topik pertanyaan seputar materi dan berdiskusi dengan rekan sekelas maupun instruktur pengampu."),
            ("Balasan Bertingkat (Threaded)", "Mendukung percakapan berbalas multi-level untuk membahas solusi teknis secara mendalam."),
            ("Moderasi Oleh Instruktur", "Trainer dapat menyematkan (pin) informasi penting atau mengunci (lock) diskusi yang sudah selesai dijawab."),
            ("Portal Berita & Pengumuman", "Feed pengumuman resmi berkategori dengan thumbnail foto menarik untuk menyampaikan agenda kegiatan lembaga.")
        ],
        image_filename="news_portal.png",
        image_caption="Tangkapan layar feed portal Berita & Pengumuman dengan kartu artikel headline, label kategori, dan waktu rilis."
    )

    # =========================================================================
    # SLIDE 14: FILOSOFI DESAIN 'WARM-LIGHT' (PENGALAMAN PENGGUNA)
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, COLOR_BG_LIGHT)
    add_header(slide14, "Standar Visual", "Filosofi Desain UI/UX: Mengapa Tampilan Kami Disukai Pengguna?")

    pillars = [
        ("1. Hangat & Bertenaga (Warm & Energetic)", "Perpaduan warna Orange (#FF6B00) dan Merah (#E11D48) membangkitkan antusiasme belajar serta memperjelas urgensi batas kuota kelas fisik."),
        ("2. Tipografi Berkarakter & Nyaman Dibaca", "Kombinasi ketegasan font Montserrat pada judul utama dengan keluwesan font Quicksand pada isi teks, menghilangkan kesan kaku."),
        ("3. Eliminasi Kontainer Abu-Abu Kusam", "Seluruh kotak abu-abu gelap (dark slate) telah diganti menjadi kartu putih bersih (#FFFFFF) berbingkai halus (#EBE5DF) di atas kanvas lembut (#FAF8F5)."),
        ("4. Responsif di Seluruh Perangkat", "Tampilan bekerja sempurna saat diakses melalui smartphone, tablet, maupun komputer desktop dengan tata letak adaptif.")
    ]
    card_w = Inches(5.7); card_h = Inches(2.55)
    for i, (title, desc) in enumerate(pillars):
        col = i % 2; row = i // 2
        cx = Inches(0.8) + col * Inches(6.0)
        cy = Inches(1.6) + row * Inches(2.8)
        c = create_card(slide14, cx, cy, card_w, card_h)
        tfc = c.text_frame; tfc.margin_left = Inches(0.3); tfc.margin_top = Inches(0.25)
        p0 = tfc.paragraphs[0]; p0.text = title; p0.font.name = "Montserrat"; p0.font.size = Pt(14); p0.font.bold = True; p0.font.color.rgb = COLOR_PRIMARY_ORANGE; p0.space_after = Pt(8)
        p1 = tfc.add_paragraph(); p1.text = desc; p1.font.name = "Quicksand"; p1.font.size = Pt(10.5); p1.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 15: CARA PENGGUNAAN SINGKAT (USER WORKFLOW GUIDE)
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15, COLOR_BG_LIGHT)
    add_header(slide15, "Panduan Operasional", "Cara Pakai Sistem: Panduan Langkah Demi Langkah Sederhana")

    steps_client = [
        ("Bagi Administrator", COLOR_CRIMSON, [
            "1. Masuk ke portal admin & kelola cabang atau instruktur.",
            "2. Buat kelas baru & tetapkan kapasitas (Offline 40 / Daring).",
            "3. Input jadwal sesi pertemuan & masukkan tautan Zoom.",
            "4. Pantau aktivitas dan keterisian kuota secara berkala."
        ]),
        ("Bagi Instruktur (Trainer)", COLOR_SKY, [
            "1. Rancang kuis dari bank soal cabang & tentukan passing grade.",
            "2. Gunakan tombol Quick Zoom saat sesi live mengajar dimulai.",
            "3. Masuk ke menu 'Verifikasi Kelulusan' untuk mengecek siswa.",
            "4. Klik 'Setujui' untuk menerbitkan sertifikat peserta."
        ]),
        ("Bagi Siswa (Peserta)", COLOR_EMERALD, [
            "1. Registrasi mandiri dengan verifikasi kode OTP email.",
            "2. Pilih kelas pelatihan pada katalog sesuai cabang & minat.",
            "3. Klik '1-Klik Masuk Zoom' & tuntaskan kuis hingga genap 20 JP.",
            "4. Unduh E-Sertifikat resmi ber-QR Code di dashboard."
        ])
    ]
    card_w = Inches(3.75); card_h = Inches(5.3)
    for i, (grp_title, grp_col, grp_items) in enumerate(steps_client):
        cx = Inches(0.8) + i * Inches(3.98)
        c = create_card(slide15, cx, Inches(1.6), card_w, card_h)
        tfc = c.text_frame; tfc.margin_left = Inches(0.3); tfc.margin_top = Inches(0.3)
        p0 = tfc.paragraphs[0]; p0.text = grp_title; p0.font.name = "Montserrat"; p0.font.size = Pt(15); p0.font.bold = True; p0.font.color.rgb = grp_col; p0.space_after = Pt(14)
        for itm in grp_items:
            pi = tfc.add_paragraph(); pi.text = itm; pi.font.name = "Quicksand"; pi.font.size = Pt(11); pi.font.color.rgb = COLOR_TEXT_MUTED; pi.space_after = Pt(10)

    # =========================================================================
    # SLIDE 16: KESIMPULAN & KESIAPAN IMPLEMENTASI (CLOSING SLIDE)
    # =========================================================================
    slide16 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide16, COLOR_BG_DARK)

    # Accent decorative glow
    glow16 = slide16.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.12))
    glow16.fill.solid(); glow16.fill.fore_color.rgb = COLOR_PRIMARY_ORANGE; glow16.line.fill.background()

    tb16 = slide16.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.5))
    tf16 = tb16.text_frame; tf16.word_wrap = True

    p0 = tf16.paragraphs[0]; p0.text = "KESIMPULAN EKSEKUTIF"; p0.font.name = "Montserrat"; p0.font.size = Pt(12); p0.font.bold = True; p0.font.color.rgb = COLOR_PRIMARY_ORANGE; p0.space_after = Pt(12)
    p1 = tf16.add_paragraph(); p1.text = "Platform Pelatihan Siap Pakai\nUntuk Pertumbuhan Lembaga Anda"; p1.font.name = "Montserrat"; p1.font.size = Pt(32); p1.font.bold = True; p1.font.color.rgb = COLOR_TEXT_WHITE; p1.space_after = Pt(16)
    
    closing_text = (
        "Sistem LMS Terpadu Multi-Cabang telah siap untuk diimplementasikan sepenuhnya:\n"
        "• Solusi Lengkap Terintegrasi: Menyatukan pendaftaran, kelas offline & daring, ujian kuis, hingga sertifikasi.\n"
        "• Standar Akreditasi 20 JP Terjamin: Kepatuhan jam belajar 900 menit tercatat akurat dan dapat diaudit.\n"
        "• Desain Modern & Bersahabat: Menghadirkan antarmuka bersih tanpa container kusam yang nyaman bagi semua pihak.\n"
        "• Kredibilitas Berskala Nasional: Otentikasi sertifikat ber-QR Code memperkuat kepercayaan mitra kerja dan publik."
    )
    p2 = tf16.add_paragraph(); p2.text = closing_text; p2.font.name = "Quicksand"; p2.font.size = Pt(13); p2.font.color.rgb = RGBColor(214, 211, 209); p2.space_after = Pt(22)
    p3 = tf16.add_paragraph(); p3.text = "Terima Kasih • Siap Mendukung Keberhasilan Transformasi Digital Pelatihan Anda"; p3.font.name = "Montserrat"; p3.font.size = Pt(12.5); p3.font.bold = True; p3.font.color.rgb = COLOR_PRIMARY_ORANGE

    # Save final client presentation
    out_file = os.path.abspath("LMS_Presentasi_Klien.pptx")
    prs.save(out_file)
    print(f"Client presentation successfully generated: {out_file}")

if __name__ == "__main__":
    build_client_presentation()
